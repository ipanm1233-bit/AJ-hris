#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Ekstrak kinerja sales dari file "08_Scheme Incentive SPV&ASR, admin, finance
Andela Agustus 2026_cair 15102026.xlsx" menjadi JSON menengah untuk
scripts/import-kinerja-sales.js.

Pemakaian:
  python scripts/export-kinerja-sales.py "<file.xlsx>" [outfile.json]

Output berisi:
  - data_blocks : blok per bulan dari sheet 'Data' (target, capaian, EC, absensi)
  - summary_sheets : ringkasan insentif dari sheet '01'..'08' (poin, POINT A-F, total cair)
  - pencairan : jadwal pencairan insentif dari sheet 'Jadwal cair'
"""

import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

MONTHS = {
    "JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4, "MEI": 5, "JUNI": 6,
    "JULI": 7, "AGUSTUS": 8, "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12,
}

# 0-based index kolom sheet 'Data' -> pemetaan header (baris "Area | SalesMan | ...")
DATA_COLUMN_MAP = {
    2: ("target", "dulux_liter"),
    3: ("target", "catylac_liter"),
    4: ("target", "maxilite_liter"),
    5: ("target", "aqs_liter"),
    6: ("target", "prima_rp"),
    7: ("target", "dcota_rp"),
    8: ("target", "ao_ici"),
    9: ("target", "ao_prima"),
    10: ("target", "ao_dcota"),
    11: ("target", "value_ici_rp"),
    12: ("target", "od_rp"),
    13: ("capaian", "dulux_liter"),
    14: ("capaian", "catylac_liter"),
    15: ("capaian", "maxilite_liter"),
    16: ("capaian", "aqs_liter"),
    17: ("capaian", "prima_rp"),
    18: ("capaian", "dcota_rp"),
    19: ("capaian", "ao_ici"),
    20: ("capaian", "ao_prima"),
    21: ("capaian", "ao_dcota"),
    22: ("capaian", "value_ici_rp"),
    23: ("capaian", "od_rp"),
    24: ("penjualan_baru", "grade"),
    25: ("effective_call", "grade"),
    26: ("absensi", "grade"),
    27: ("penjualan_baru", "outlet_baru"),
    28: ("effective_call", "target"),
    29: ("effective_call", "realisasi"),
    30: ("effective_call", "persen"),
    31: ("absensi", "hadir"),
    32: ("absensi", "absen"),
    33: ("absensi", "sakit"),
    34: ("absensi", "cuti"),
    35: ("absensi", "alpa"),
    36: ("absensi", "persen_hadir"),
    37: ("absensi", "hari_kerja"),
    38: ("absensi", "um_luar_kota"),
    39: ("absensi", "day"),
    40: ("absensi", "diff_absen"),
}

SUMMARY_EXTRA_HEADERS = {
    "UM Luar Kota 2025": "um_luar_kota",
    "UM Luar Kota 2026": "um_luar_kota",
    "Revisi Insentif AO Jan 2026": "revisi_insentif_ao",
    "Q1 2026": "bonus_q1",
    "Q2": "bonus_q2",
    "Adjutsment Juni 2026": "adjustment_juni",
    "Adjustment Juni 2026": "adjustment_juni",
}


def clean_text(value):
    s = "" if value is None else str(value)
    s = "".join(ch for ch in s if ch.isprintable())
    return re.sub(r"\s+", " ", s).strip()


def as_number(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            return None
        return round(float(value), 6)
    s = clean_text(value)
    if not s:
        return None
    if re.match(r"^#.*->$|^(#DIV/0!|#N/A|#VALUE!|#REF!)", s):
        return None
    if re.match(r"^-?\d+(\.\d+)?%$", s):
        return round(float(s.rstrip("%")) / 100, 6)
    s2 = s.replace("Rp", "").replace(".", "").replace(",", ".").strip()
    try:
        return round(float(s2), 6)
    except ValueError:
        return None


def cell(value):
    """Nilai sel mentah: angka bila numerik, teks bila grade/label."""
    number = as_number(value)
    if number is not None:
        return number
    text = clean_text(value)
    return text or None


def total_rows(ws):
    return list(ws.iter_rows(values_only=True))


def extract_data_sheet(ws):
    rows = total_rows(ws)
    out_blocks = []
    tahun = 2026
    for r in rows:
        label = clean_text(r[0] if r else "")
        if label.startswith("TAHUN"):
            match = re.search(r"\d{4}", label)
            if match:
                tahun = int(match.group(0))
    i = 0
    while i < len(rows):
        label = clean_text(rows[i][0] if rows[i] else "").upper()
        if label in MONTHS:
            header_idx = i + 1
            header = header_idx if header_idx < len(rows) else None
            if not header or clean_text(rows[header][0] if rows[header] else "").upper() != "AREA":
                i += 1
                continue
            j = header_idx + 1
            data_rows = []
            while j < len(rows):
                area = clean_text(rows[j][0] if rows[j] else "")
                alias = clean_text(rows[j][1] if len(rows[j]) > 1 and rows[j][1] else "")
                if not area:
                    break
                row_out = {"area": area, "alias": alias, "groups": {}}
                for idx, (group, field) in DATA_COLUMN_MAP.items():
                    val = cell(rows[j][idx]) if idx < len(rows[j]) else None
                    row_out["groups"].setdefault(group, {})[field] = val
                data_rows.append(row_out)
                j += 1
            out_blocks.append({
                "bulan": clean_text(rows[i][0]).upper(),
                "periode": f"{tahun:04d}-{MONTHS[label]:02d}",
                "rows": data_rows,
            })
            i = j
        else:
            i += 1
    return tahun, out_blocks


def extract_summary_sheets(wb):
    hasil = []
    for sheet in wb.sheetnames:
        if not re.fullmatch(r"0?\d{1,2}", sheet):
            continue
        month_num = int(sheet) if sheet.isdigit() else int(sheet.zfill(2))
        if not (1 <= month_num <= 12):
            continue
        ws = wb[sheet]
        rows = total_rows(ws)
        header_idx = None
        for idx, r in enumerate(rows):
            vals = [clean_text(c).upper() for c in r if c is not None]
            if "AREA" in vals and any(v.startswith("ASR") for v in vals):
                header_idx = idx
                break
        if header_idx is None:
            continue
        header = list(rows[header_idx])
        header_prev = list(rows[header_idx - 1]) if header_idx > 0 else []
        # map kolom berdasarkan teks header (baris utama + baris jabatan di atasnya
        # untuk 'insentif @ Admin/Finance')
        col_map = {}
        for idx, c in enumerate(header):
            label = clean_text(c).upper()
            if label == "ASR":
                col_map["alias"] = idx
            elif label == "POIN":
                col_map["poin"] = idx
            elif label.startswith("POINT A"):
                col_map["point_a"] = idx
            elif label.startswith("POINT B"):
                col_map["point_b"] = idx
            elif label.startswith("POINT C"):
                col_map["point_c"] = idx
            elif label.startswith("POINT D"):
                col_map["point_d"] = idx
            elif label.startswith("POINT E"):
                col_map["point_e"] = idx
            elif label.startswith("POINT F"):
                col_map["point_f"] = idx
            elif label == "TOTAL":
                col_map["total"] = idx
            elif label in SUMMARY_EXTRA_HEADERS:
                col_map[SUMMARY_EXTRA_HEADERS[label]] = idx
        for idx, c in enumerate(header_prev + [None] * (len(header) - len(header_prev))):
            label = clean_text(c).upper() if c is not None else ""
            if not label:
                continue
            if label.startswith("INSENTIF @ ADMIN") and "insentif_admin" not in col_map:
                col_map["insentif_admin"] = idx if idx < len(header) else len(header) - 1
            elif (label.startswith("INSENTIF @FINANCE") or label.startswith("INSENTIF @ FINANCE")) and "insentif_finance" not in col_map:
                col_map["insentif_finance"] = idx if idx < len(header) else len(header) - 1
        data_rows = []
        j = header_idx + 1
        while j < len(rows):
            area = clean_text(rows[j][0] if rows[j] else "")
            if not area:
                j += 1
                continue
            vals = [clean_text(c).upper() for c in rows[j] if c is not None]
            if vals and vals[0].startswith("SALESMAN"):
                j += 1
                continue
            row_out = {"area": area}
            for field, idx in col_map.items():
                row_out[field] = cell(rows[j][idx]) if idx < len(rows[j]) else None
            data_rows.append(row_out)
            j += 1
            if j < len(rows) and clean_text(rows[j][0] if rows[j] else "") == "":
                # berhenti setelah baris data berakhir (dua baris kosong berurutan)
                following = [clean_text(r[0]) for r in rows[j:j + 2] if r and r[0] is not None]
                if not any(following):
                    break
        hasil.append({"sheet": sheet, "bulan_num": month_num, "rows": data_rows})
    return hasil


def extract_pencairan(ws):
    rows = total_rows(ws)
    hasil = []
    rencana = False
    for r in rows:
        a = clean_text(r[0] if r else "")
        if a.upper().startswith("PERIODE PENJUALAN"):
            rencana = True
            continue
        if a.upper().startswith("PENCAIRAN INSENTIF"):
            rencana = False
        if rencana and a and re.search(r"\d{4}", a):
            hasil.append({
                "bulan_achiemenent": a,
                "tgl_cair": clean_text(r[1] if len(r) > 1 and r[1] else ""),
                "dihitung_dari": clean_text(r[2] if len(r) > 2 and r[2] else ""),
            })
    return hasil


def main():
    if len(sys.argv) < 2:
        print("Pemakaian: python export-kinerja-sales.py <file.xlsx> [outfile.json]")
        return 2
    src = os.path.abspath(sys.argv[1])
    out = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else src + ".extract.json"
    wb = openpyxl.load_workbook(src, data_only=True)
    tahun, blocks = extract_data_sheet(wb["Data"])
    summaries = extract_summary_sheets(wb)
    pencairan = extract_pencairan(wb["Jadwal cair"]) if "Jadwal cair" in wb.sheetnames else []

    blob = {
        "source_file": os.path.basename(src),
        "tahun": tahun,
        "data_blocks": blocks,
        "summary_sheets": summaries,
        "pencairan": pencairan,
    }
    with open(out, "w", encoding="utf8") as fh:
        json.dump(blob, fh, ensure_ascii=False, indent=1)
    print(f"OK: {out}")
    print(f"  blok bulanan : {len(blocks)} ({', '.join(b['periode'] for b in blocks)})")
    print(f"  sheet insentif: {len(summaries)} ({', '.join(s['sheet'] for s in summaries)})")
    print(f"  jadwal cair  : {len(pencairan)}")
    rows_total = sum(len(b["rows"]) for b in blocks)
    print(f"  total baris  : {rows_total}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
